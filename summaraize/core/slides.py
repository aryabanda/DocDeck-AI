"""Slide content (extractive or LLM-written) and PowerPoint generation."""
import io
import os

from pptx import Presentation
from pptx.util import Pt

from core.llm import chat, parse_json

MAX_CLUSTER_CHARS = 6000  # keep each LLM request small and cheap

SLIDE_SYSTEM = (
    "You write presentation slides. Use ONLY the text the user provides; never add "
    "facts, numbers or names that are not in it. Return ONLY a JSON object, no prose, "
    'in exactly this shape: {"title": "<max 8 words>", "bullets": ["...", "...", "..."]}. '
    "Exactly 3 bullets, each at most 20 words."
)


def extractive_slide_specs(topics: dict, keywords: dict, per_slide: int = 5) -> list[dict]:
    """Your original behaviour: up to 5 sentences per slide, titled by topic keywords."""
    specs = []
    for label, sentences in topics.items():
        for i in range(0, len(sentences), per_slide):
            title = f"Topic: {keywords[label]}" + (" (continued)" if i else "")
            specs.append({"title": title, "bullets": sentences[i : i + per_slide]})
    return specs


def _valid(spec) -> bool:
    return (
        isinstance(spec, dict)
        and isinstance(spec.get("title"), str)
        and isinstance(spec.get("bullets"), list)
        and len(spec["bullets"]) > 0
        and all(isinstance(b, str) for b in spec["bullets"])
    )


def llm_slide_spec(sentences: list[str], llm=chat) -> dict | None:
    """Ask the LLM for one slide from one topic cluster.

    Bad/invalid model output is retried once, then None is returned.
    Errors from the LLM call itself (missing key, auth, network, rate limit)
    are NOT swallowed: they propagate so the user sees the real problem.
    """
    source = " ".join(sentences)[:MAX_CLUSTER_CHARS]
    messages = [
        {"role": "system", "content": SLIDE_SYSTEM},
        {"role": "user", "content": f"Text:\n{source}"},
    ]
    for _ in range(2):
        raw = llm(messages)
        try:
            spec = parse_json(raw)
        except ValueError:  # json.JSONDecodeError is a ValueError
            spec = None
        if _valid(spec):
            return {"title": spec["title"].strip(), "bullets": [b.strip() for b in spec["bullets"][:3]]}
        messages.append({"role": "assistant", "content": raw})
        messages.append({"role": "user", "content": "That was not valid. Reply with ONLY the JSON object."})
    return None


def llm_slide_specs(topics: dict, keywords: dict, llm=chat) -> list[dict]:
    """One LLM call per topic cluster. If the model keeps returning unusable output
    for a cluster, that slide falls back to extractive text and is flagged."""
    specs = []
    for label, sentences in topics.items():
        spec = llm_slide_spec(sentences, llm=llm)
        if spec is None:
            spec = {"title": f"Topic: {keywords[label]}", "bullets": sentences[:5], "fallback": True}
        specs.append(spec)
    return specs


def _clear_slides(prs: Presentation) -> None:
    """Remove any slides that ship inside a template file."""
    id_list = prs.slides._sldIdLst
    for slide_id in list(id_list):
        prs.part.drop_rel(slide_id.rId)
        id_list.remove(slide_id)


def build_presentation(specs: list[dict], title: str, subtitle: str = "", template: str | None = None) -> bytes:
    """Build a .pptx and return its bytes (no temp files, safe for web apps)."""
    if template and os.path.exists(template):
        prs = Presentation(template)
        _clear_slides(prs)
    else:
        prs = Presentation()

    cover = prs.slides.add_slide(prs.slide_layouts[0])
    cover.shapes.title.text = title or "Summary"
    cover.placeholders[1].text = subtitle

    for spec in specs:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        size = Pt(20) if len(spec["bullets"]) <= 3 else Pt(16)
        for i, bullet in enumerate(spec["bullets"]):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.text = bullet
            for run in paragraph.runs:
                run.font.size = size

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
