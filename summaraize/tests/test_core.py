"""Offline tests: a fake embedder and a fake LLM stand in for the real ones,
so the whole pipeline is tested without an API key or model downloads."""
import hashlib
import io
import json
import os

import numpy as np
import pytest
from pptx import Presentation

from core.extractive import divide_into_topics, extractive_sentences
from core.llm import parse_json
from core.pdf_utils import extract_pages, full_text
from core.rag import NOT_FOUND, answer, build_index, chunk_pages
from core.slides import (build_presentation, extractive_slide_specs,
                         llm_slide_spec, llm_slide_specs)

SAMPLE = os.path.join(os.path.dirname(__file__), "..", "docs", "sample.pdf")


class FakeEmbedder:
    """Deterministic bag-of-words hashing embedder (keyword overlap = similarity)."""

    def encode(self, texts):
        out = []
        for t in texts:
            v = np.zeros(256)
            for w in t.lower().split():
                v[int(hashlib.md5(w.strip(".,;:()").encode()).hexdigest(), 16) % 256] += 1
            n = np.linalg.norm(v)
            out.append((v / n if n else v).tolist())
        return out


def test_parse_json_handles_fences_and_chatter():
    assert parse_json('```json\n{"a": 1}\n```') == {"a": 1}
    assert parse_json('Sure! Here you go: {"a": 1} hope it helps') == {"a": 1}
    with pytest.raises(Exception):
        parse_json("no json here")


def test_chunking_keeps_pages_and_overlap():
    pages = [{"page": 2, "text": " ".join(f"w{i}" for i in range(100))},
             {"page": 5, "text": "short page text"}]
    chunks = chunk_pages(pages, chunk_words=40, overlap=10)
    assert {c["page"] for c in chunks} == {2, 5}
    assert len({c["id"] for c in chunks}) == len(chunks)          # unique ids
    first, second = chunks[0]["text"].split(), chunks[1]["text"].split()
    assert first[-10:] == second[:10]                              # overlap works
    assert " ".join(c["text"] for c in chunks if c["page"] == 5) == "short page text"
    with pytest.raises(ValueError):
        chunk_pages(pages, chunk_words=10, overlap=10)


def test_retrieval_and_answer_cite_correct_page():
    pages = [
        {"page": 1, "text": "The cat sat on the mat and purred quietly all afternoon."},
        {"page": 2, "text": "Quarterly revenue grew twelve percent driven by subscription renewals."},
        {"page": 3, "text": "The hiking trail closes in winter because of heavy snowfall."},
    ]
    index = build_index(pages, FakeEmbedder(), chunk_words=30, overlap=5)
    hits = index.retrieve("what happened to subscription revenue", k=2)
    assert hits[0]["page"] == 2

    seen = {}
    def fake_llm(messages):
        seen["prompt"] = messages[-1]["content"]
        return "Revenue grew twelve percent [p. 2]."
    result = answer("what happened to subscription revenue", index, k=2, llm=fake_llm)
    assert "[p. 2]" in seen["prompt"] and "[p. 2]" in result["answer"]
    assert result["sources"][0]["page"] == 2


def test_llm_slide_spec_retries_then_falls_back():
    calls = {"n": 0}
    def flaky(messages):
        calls["n"] += 1
        return "garbage" if calls["n"] == 1 else '{"title": "T", "bullets": ["a", "b", "c", "d"]}'
    spec = llm_slide_spec(["Some sentence about things."], llm=flaky)
    assert spec == {"title": "T", "bullets": ["a", "b", "c"]} and calls["n"] == 2

    assert llm_slide_spec(["x y z w."], llm=lambda m: "never json") is None
    topics, kw = {0: ["Alpha beta gamma delta."]}, {0: "alpha"}
    fallback = llm_slide_specs(topics, kw, llm=lambda m: "never json")
    assert fallback[0]["bullets"] == ["Alpha beta gamma delta."]


def test_build_presentation_is_a_valid_pptx():
    specs = [{"title": "One", "bullets": ["x", "y", "z"]}, {"title": "Two", "bullets": ["q"]}]
    data = build_presentation(specs, "My deck", "sub")
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 3                                    # cover + 2
    assert prs.slides[1].shapes.title.text == "One"


def test_end_to_end_on_sample_pdf():
    pages = extract_pages(SAMPLE)
    assert pages and all("page" in p and p["text"] for p in pages)
    sentences = extractive_sentences(full_text(pages), ratio=0.2)
    assert sentences
    topics, keywords = divide_into_topics(sentences, n_topics=3)
    assert set(topics) == set(keywords)                            # keywords match clusters
    assert sum(len(v) for v in topics.values()) == len(sentences)
    specs = extractive_slide_specs(topics, keywords)
    assert Presentation(io.BytesIO(build_presentation(specs, "t"))).slides
    index = build_index(pages, FakeEmbedder())
    assert answer("test question", index, llm=lambda m: NOT_FOUND)["answer"] == NOT_FOUND


def test_llm_call_errors_are_not_swallowed():
    def broken(messages):
        raise RuntimeError("LLM_API_KEY is not set")
    with pytest.raises(RuntimeError):
        llm_slide_spec(["Some sentence about things."], llm=broken)
    fallback = llm_slide_specs({0: ["Alpha beta gamma delta."]}, {0: "alpha"}, llm=lambda m: "never json")
    assert fallback[0].get("fallback") is True


def test_retrieval_metrics_and_summary_scoring():
    from eval.eval_retrieval import evaluate
    from eval.eval_summaries import deck_text
    pages = [
        {"page": 1, "text": "The cat sat on the mat and purred quietly all afternoon."},
        {"page": 2, "text": "Quarterly revenue grew twelve percent driven by subscription renewals."},
        {"page": 3, "text": "The hiking trail closes in winter because of heavy snowfall."},
    ]
    qs = [
        {"question": "how much did revenue grow", "page": 2, "evidence": "grew twelve percent"},
        {"question": "hiking trail winter snowfall", "page": 3},
        {"question": "completely unrelated zebra", "page": 1, "evidence": "phrase that is not there"},
    ]
    r = evaluate(pages, qs, FakeEmbedder(), chunk_words=30, overlap=5, k=3)
    assert r["hit@1"] == pytest.approx(2 / 3) and r["hit@3"] == pytest.approx(2 / 3)
    assert 0 < r["mrr"] <= 1
    assert deck_text([{"bullets": ["a", "b"]}, {"bullets": ["c"]}]) == "a b c"
